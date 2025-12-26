"""
Tool Registry - Maps tool names to APIs and AI services
"""
import httpx
from google import genai
from google.genai import types
from typing import Dict, Any
import logging
import os
import uuid
from datetime import datetime
from io import BytesIO
from docx import Document

logger = logging.getLogger(__name__)

# External API base URL for image generation
EXTERNAL_API_BASE_URL = "http://13.205.132.169:8000"

# Configure Gemini API
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if GEMINI_API_KEY:
    client = genai.Client(api_key=GEMINI_API_KEY)
else:
    client = None
    logger.warning("GEMINI_API_KEY not found in environment variables")

class ToolRegistry:
    """Registry of available tools and their execution handlers"""
    
    def __init__(self):
        self.tools = {
            "blog_generator": self._execute_blog_generator,
            "image_generator": self._execute_image_generator,
            "caption_generator": self._execute_caption_generator,
            "content_optimizer": self._execute_content_optimizer,
            "hashtag_generator": self._execute_hashtag_generator,
            "post_to_x": self._execute_post_to_x,
        }
        
        # Initialize Gemini client
        self.client = client
        if not client:
            logger.warning("Gemini client not initialized - API key missing")
        
        # Import S3 service
        from app.services.s3_service import s3_service
        self.s3_service = s3_service
    
    async def execute_tool(self, tool_name: str, config: Dict[str, Any]) -> Dict[str, Any]:
        """Execute a tool with given configuration"""
        if tool_name not in self.tools:
            raise ValueError(f"Unknown tool: {tool_name}")
        
        handler = self.tools[tool_name]
        return await handler(config)
    
    async def _execute_blog_generator(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """Generate blog post using Gemini AI and save to S3"""
        logger.info(f"Generating blog with config: {config}")
        
        if not self.client:
            raise ValueError("Gemini API not configured - add GEMINI_API_KEY to .env")
        
        try:
            topic = config.get('topic', 'general topic')
            word_count = config.get('word_count', 1500)
            
            # Generate blog content using Gemini
            prompt = f"""Write a comprehensive, well-structured blog post about: {topic}

Requirements:
- Target length: {word_count} words
- Include an engaging introduction
- Use clear headings and subheadings
- Provide detailed explanations and examples
- Include a strong conclusion
- Write in a professional yet accessible tone
- Make it informative and valuable to readers

Generate the complete blog post:"""

            response = self.client.models.generate_content(
                model='gemini-2.0-flash-exp',
                contents=prompt
            )
            
            blog_content = response.text.strip()
            
            # Create DOCX document
            doc = Document()
            doc.add_heading(topic, 0)
            
            # Add content paragraphs
            for paragraph in blog_content.split('\n\n'):
                if paragraph.strip():
                    if paragraph.startswith('#'):
                        # Handle markdown headings
                        heading_text = paragraph.lstrip('#').strip()
                        doc.add_heading(heading_text, level=1)
                    else:
                        doc.add_paragraph(paragraph.strip())
            
            # Save to buffer
            docx_buffer = BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)
            
            # Upload to S3
            file_id = str(uuid.uuid4())
            docx_key = f"blogs/docs/blog_{file_id}.docx"
            docx_url = await self.s3_service.upload_buffer(
                docx_key,
                docx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
            logger.info("Blog generated and uploaded successfully")
            return {
                "status": "success",
                "tool": "blog_generator",
                "result": {
                    "id": file_id,
                    "topic": topic,
                    "docx_url": docx_url,
                    "word_count": len(blog_content.split()),
                    "created_at": datetime.utcnow().isoformat()
                }
            }
            
        except Exception as e:
            logger.error(f"Blog generation failed: {str(e)}")
            raise
    
    async def _execute_image_generator(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """Generate image using external API"""
        logger.info(f"Generating image with config: {config}")
        
        try:
            topic = config.get('topic', 'inspirational scene')
            
            # Call external API for image generation
            async with httpx.AsyncClient(timeout=120.0) as http_client:
                response = await http_client.post(
                    f"{EXTERNAL_API_BASE_URL}/api/v1/generate/motivational_post",
                    json={"topic": topic}
                )
                response.raise_for_status()
                result = response.json()
                
                logger.info("Image generated successfully via external API")
                return {
                    "status": "success",
                    "tool": "image_generator",
                    "result": {
                        "id": result.get("id", str(uuid.uuid4())),
                        "topic": topic,
                        "image_url": result.get("image_url", ""),
                        "quote_text": result.get("quote_text", ""),
                        "created_at": result.get("created_at", datetime.utcnow().isoformat())
                    }
                }
            
        except Exception as e:
            logger.error(f"Image generation failed: {str(e)}")
            raise
    
    async def _execute_caption_generator(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """Generate social media caption using Gemini AI"""
        logger.info(f"Generating caption with config: {config}")
        
        if not self.client:
            raise ValueError("Gemini API not configured - add GEMINI_API_KEY to .env")
        
        try:
            # Extract parameters
            topic = config.get('topic', 'general content')
            platform = config.get('platform', 'linkedin')
            tone = config.get('tone', 'professional')
            max_length = config.get('max_length', 200)
            include_emojis = config.get('include_emojis', True)
            include_hashtags = config.get('include_hashtags', True)
            
            # Build prompt
            prompt = f"""Generate an engaging {platform} caption about: {topic}

Requirements:
- Tone: {tone}
- Maximum length: {max_length} characters
- Include emojis: {include_emojis}
- Include hashtags: {include_hashtags}
- Make it compelling and shareable
- Use proper formatting

Generate only the caption, no additional text."""

            # Call Gemini
            response = self.client.models.generate_content(
                model='gemini-2.0-flash-exp',
                contents=prompt
            )
            caption = response.text.strip()
            
            logger.info(f"Caption generated: {len(caption)} characters")
            
            return {
                "status": "success",
                "tool": "caption_generator",
                "result": {
                    "caption": caption,
                    "length": len(caption),
                    "platform": platform,
                    "tone": tone
                }
            }
            
        except Exception as e:
            logger.error(f"Caption generation failed: {str(e)}")
            raise
    
    async def _execute_content_optimizer(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """Optimize content using Gemini AI"""
        logger.info(f"Optimizing content with config: {config}")
        
        if not self.client:
            raise ValueError("Gemini API not configured - add GEMINI_API_KEY to .env")
        
        try:
            # Extract parameters
            content = config.get('content', '')
            optimization_goal = config.get('goal', 'engagement')
            target_audience = config.get('audience', 'professionals')
            platform = config.get('platform', 'linkedin')
            
            if not content:
                raise ValueError("Content is required for optimization")
            
            # Build prompt
            prompt = f"""Optimize the following content for {platform}:

Original Content:
{content}

Optimization Goals:
- Goal: {optimization_goal}
- Target Audience: {target_audience}
- Platform: {platform}

Please provide:
1. Optimized version of the content
2. Key improvements made
3. SEO/engagement tips

Format as JSON with keys: optimized_content, improvements, tips"""

            # Call Gemini
            response = self.client.models.generate_content(
                model='gemini-2.0-flash-exp',
                contents=prompt
            )
            result_text = response.text.strip()
            
            logger.info("Content optimized successfully")
            
            return {
                "status": "success",
                "tool": "content_optimizer",
                "result": {
                    "original": content,
                    "optimized": result_text,
                    "goal": optimization_goal,
                    "platform": platform
                }
            }
            
        except Exception as e:
            logger.error(f"Content optimization failed: {str(e)}")
            raise
    
    async def _execute_hashtag_generator(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """Generate relevant hashtags using Gemini AI"""
        logger.info(f"Generating hashtags with config: {config}")
        
        try:
            # Extract parameters
            topic = config.get('topic', '')
            content = config.get('content', '')
            count = config.get('count', 5)
            platform = config.get('platform', 'linkedin')
            category = config.get('category', 'general')
            
            if not topic and not content:
                raise ValueError("Either topic or content is required")
            
            # Build prompt
            context = content if content else topic
            prompt = f"""Generate {count} relevant hashtags for {platform} based on:

Content/Topic: {context}
Category: {category}

Requirements:
- Hashtags should be popular and discoverable
- Mix of broad and niche hashtags
- Appropriate for {platform}
- No spaces, use camelCase or lowercase
- Return as a comma-separated list

Example format: #AI, #TechInnovation, #FutureOfWork"""

            # Call Gemini
            response = self.client.models.generate_content(
                model='gemini-2.5-flash',
                contents=prompt
            )
            hashtags_text = response.text.strip()
            
            # Parse hashtags
            hashtags = [tag.strip() for tag in hashtags_text.replace('\n', ',').split(',') if tag.strip()]
            hashtags = [tag if tag.startswith('#') else f'#{tag}' for tag in hashtags][:count]
            
            logger.info(f"Generated {len(hashtags)} hashtags")
            
            return {
                "status": "success",
                "tool": "hashtag_generator",
                "result": {
                    "hashtags": hashtags,
                    "count": len(hashtags),
                    "platform": platform,
                    "category": category
                }
            }
            
        except Exception as e:
            logger.error(f"Hashtag generation failed: {str(e)}")
            raise
    
    async def _execute_post_to_x(self, config: Dict[str, Any]) -> Dict[str, Any]:
        """
        Prepare content for X (Twitter) posting (MCP tool - does NOT post directly)
        
        This tool:
        1. Validates that content exists and is approved
        2. Prepares the final posting payload
        3. Creates a posting_job record with status='ready'
        4. Returns the job_id and prepared payload
        
        It does NOT:
        - Call X API
        - Perform browser automation
        - Schedule posts
        
        The user must manually post using the prepared payload.
        """
        logger.info(f"Preparing X post with config: {config}")
        
        try:
            # Import database and models
            from app.database import supabase
            from app.models import CONTENT_TABLE, POSTING_JOBS_TABLE
            
            # Extract parameters
            content_id = config.get('content_id')
            workspace_id = config.get('workspace_id')
            
            if not content_id or not workspace_id:
                raise ValueError("content_id and workspace_id are required")
            
            # STEP 1: Validate content exists and is approved
            content_response = supabase.table(CONTENT_TABLE)\
                .select("*")\
                .eq("id", content_id)\
                .eq("workspace_id", workspace_id)\
                .is_("deleted_at", "null")\
                .execute()
            
            if not content_response.data:
                raise ValueError(f"Content {content_id} not found or deleted")
            
            content = content_response.data[0]
            
            if content['status'] != 'approved':
                raise ValueError(f"Content must be approved before posting. Current status: {content['status']}")
            
            # STEP 2: Check if posting_job already exists for this content
            existing_job_response = supabase.table(POSTING_JOBS_TABLE)\
                .select("*")\
                .eq("content_id", content_id)\
                .execute()
            
            if existing_job_response.data:
                existing_job = existing_job_response.data[0]
                logger.info(f"Posting job already exists: {existing_job['id']}")
                return {
                    "status": "success",
                    "tool": "post_to_x",
                    "result": {
                        "job_id": existing_job['id'],
                        "content_id": content_id,
                        "status": existing_job['status'],
                        "prepared_payload": existing_job['prepared_payload'],
                        "message": "Posting job already exists"
                    }
                }
            
            # STEP 3: Prepare final payload
            content_data = content['data']
            content_type = content['content_type']
            
            # Build payload based on content type
            prepared_payload = {
                "content_id": content_id,
                "platform": "x",
                "created_at": datetime.utcnow().isoformat()
            }
            
            if content_type == 'caption':
                prepared_payload['post_text'] = content_data.get('caption', '')
            elif content_type == 'optimized_content':
                prepared_payload['post_text'] = content_data.get('optimized', content_data.get('content', ''))
            elif content_type == 'blog_post':
                # For blog posts, create a teaser
                title = content.get('title', 'New Blog Post')
                docx_url = content_data.get('docx_url', '')
                prepared_payload['post_text'] = f"📝 {title}\n\nRead the full article: {docx_url}"
            else:
                # Generic fallback
                prepared_payload['post_text'] = str(content_data.get('text', content_data.get('content', '')))
            
            # Add hashtags if available
            if content_type == 'hashtags':
                prepared_payload['hashtags'] = content_data.get('hashtags', [])
            
            # Add image URL if available
            if 'image_url' in content_data:
                prepared_payload['image_url'] = content_data['image_url']
            if 'cover_url' in content_data:
                prepared_payload['image_url'] = content_data['cover_url']
            
            # X-specific formatting hints (280 char limit for tweets)
            prepared_payload['formatting_hints'] = {
                "max_length": 280,
                "supports_markdown": False,
                "supports_images": True,
                "supports_videos": True
            }
            
            # STEP 4: Create posting_job record
            posting_job_data = {
                "workspace_id": workspace_id,
                "content_id": content_id,
                "platform": "x",
                "status": "ready",
                "prepared_payload": prepared_payload,
                "retry_count": 0
            }
            
            job_response = supabase.table(POSTING_JOBS_TABLE)\
                .insert(posting_job_data)\
                .execute()
            
            if not job_response.data:
                raise ValueError("Failed to create posting job")
            
            created_job = job_response.data[0]
            logger.info(f"Created posting job: {created_job['id']}")
            
            # STEP 5: Return success response
            return {
                "status": "success",
                "tool": "post_to_x",
                "result": {
                    "job_id": created_job['id'],
                    "content_id": content_id,
                    "status": "ready",
                    "prepared_payload": prepared_payload,
                    "message": "Content prepared for X posting. User must manually post."
                }
            }
            
        except Exception as e:
            logger.error(f"X posting preparation failed: {str(e)}")
            raise
    
    def is_tool_available(self, tool_name: str) -> bool:
        """Check if a tool is registered"""
        return tool_name in self.tools
    
    def list_tools(self) -> list:
        """Get list of all available tools"""
        return list(self.tools.keys())


# Singleton instance
tool_registry = ToolRegistry()